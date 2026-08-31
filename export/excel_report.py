# PRISAnalyzer.Class.ExcelReport 

from dataclasses import dataclass, field
from typing import Optional
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.cell.cell import MergedCell


@dataclass
class ECell:
    text: str
    font_size: int = 11
    font_name : str = "標楷體"
    bold: bool = False
    is_formula: bool = False
    bg_color: Optional[str] = None # openpyxl 用色碼字串(如 "FFFF00"),不是 XLColor 物件
    vertical_align: Optional[str] = None # openpyxl 用字串("center", "top", "bottom")
    horizontal_align: Optional[str] = None # openpyxl 用字串("left", "center", "right")
    merge_row_range: Optional[int] = None
    merge_col_range: Optional[int] = None
    hyperlink : Optional[str] = None # openpyxl 用字串(目標位置或 URL),不是 XLHyperlink 物件

## Sheet裡面的表格的Header
@dataclass
class ExcelColumnHeader:
    name : str = ""
    vertical_align: Optional[str] = None # openpyxl 用字串("center", "top", "bottom")
    horizontal_align: Optional[str] = None # openpyxl 用字串("left", "center", "right")
    sub_headers: list["ExcelColumnHeader"] = field(default_factory=list)
    wrap_text: bool = False
    # public (int MaxCharNum, double HeightPerRow)? ForceHeightCharNumPerRow { get; set; }
    max_char : Optional[int] = None #ForceHeightCharNumPerRow
    height_per_row: Optional[float] = None #ForceHeightCharNumPerRow
    
    merge_row_range : Optional[int] = None
    width:int = -1
    height: int = -1
    is_formula: bool = False

    def get_max_depth(self) -> int:
        if not self.sub_headers:
            #print("no sub header!!!")
            return 1

        sun_max_d = 0
        for col in self.sub_headers:
            tmp_d = col.get_max_depth()
            if tmp_d > sun_max_d:
                sun_max_d = tmp_d

        return sun_max_d+1

    def get_span(self) -> int:
        ret = 0
        if not self.sub_headers:
            #debug
            #print("no sub header!!!")
            return 1

        for col in self.sub_headers:
            ret += col.get_span()

        return ret


@dataclass
class ExcelReport():
    headers : list[ExcelColumnHeader] = field(default_factory=list)
    rows : list[list[ECell]] = field(default_factory=list)

    def export(self,wb : Workbook, new_tab_name : str):
        ws = wb.create_sheet(new_tab_name)
        x = 1
        max_d = float("-inf")
        for col in self.headers:
            self.render_col_header(ws, col, x, 1)
            x += col.get_span()
            d = col.get_max_depth()
            if max_d < d:
                max_d = d

        self.render_rows(ws, self.rows, max_d+1)

    def render_col_header(self, ws:Worksheet, header:ExcelColumnHeader, x :int, d:int):
        cell = ws.cell(d,x)
        cell.font = Font(size=11, name="標楷體")

        horizontal = header.horizontal_align if header.horizontal_align is not None else "center"
        vertical = header.vertical_align if header.vertical_align is not None else "center"

        cell.alignment = Alignment(horizontal = horizontal, vertical = vertical, wrap_text=header.wrap_text)

        if header.width != -1:
            ws.column_dimensions[cell.column_letter].width = header.width
        if header.height != -1:
            #ws.row_dimensions[cell.column_letter].height = header.height
            ws.row_dimensions[d].height = header.height
        if header.is_formula:
            cell.value = f"={header.name}"
        else:
            cell.value = header.name

        span = header.get_span()
        # C# 對照 ws.Range(d, x, d + ((header.MergeRowRange) ?? 1) - 1, x + span - 1).Merge();
        # firstcellrow, firstcellcol, lastcellrow, lastcellcol
        end_row = d + (header.merge_row_range or 1) -1
        end_col = x + span -1
        ws.merge_cells(start_row=d, start_column=x, end_row= end_row, end_column= end_col)

        if header.height_per_row is not None and header.max_char is not None:
            import math
            row_num = math.ceil(len(header.name) / header.max_char)
            height = row_num * header.height_per_row
            current_height = ws.row_dimensions[d].height or 0
            if current_height < height:
                ws.row_dimensions[d].height = height

        sub_x = x
        for sub_col in header.sub_headers:
            self.render_col_header(ws, sub_col, sub_x, d+1)
            sub_x += sub_col.get_span()

    def render_rows(self, ws : Worksheet, rows: list[list[ECell]], start_row_num : int):
        from common.util import convert_type
        r = start_row_num
        for row in rows:
            c = 1
            for col in row:
                cell = ws.cell(r,c)

                # 針對超連結做修改 - python需要加
                font_color = "0000FF" if col.hyperlink is not None else None
                font_underline = "single" if col.hyperlink is not None else None

                cell.font = Font(
                    bold=col.bold, 
                    size=col.font_size, 
                    name= col.font_name,
                    color=font_color,
                    underline=font_underline
                )

                horizon = col.horizontal_align if col.horizontal_align is not None else "center"
                vertical = col.vertical_align if col.vertical_align is not None else "center"
                cell.alignment = Alignment(horizontal=horizon, vertical= vertical, wrap_text=True)


                if col.is_formula and col.text: #如果是空字串就不用轉換成公式了
                    cell.value = f"={col.text}"
                elif isinstance(cell, MergedCell) == False:
                    cell.value = convert_type(col.text)
                    #cell.value = col.text


                if col.merge_row_range is not None or col.merge_col_range is not None:
                    end_row = r + (col.merge_row_range or 1) - 1
                    end_col = c + (col.merge_col_range or 1) - 1
                    ws.merge_cells(start_row=r,start_column=c,end_row=end_row,end_column=end_col)

                if col.bg_color is not None:
                    cell.fill = PatternFill(
                        start_color=col.bg_color,
                        end_color=col.bg_color,
                        fill_type="solid"
                    )

                if col.hyperlink is not None:
                    cell.hyperlink = col.hyperlink

                c += 1
            r += 1

