## PRISAnalyzer.PresidentReportESI.Class.Analysis.Templates.Template
from dataclasses import dataclass, field
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter
from abc import ABC, abstractmethod
from enum import Enum, auto

#config
from config import *

# PRIS
from export.excel_report import *

class TemplateType(Enum):
    E1 = 1
    E2 = auto()
    E3 = auto()
    E4 = auto()
    E5 = auto()
    E6 = auto()
    E7 = auto()
    F1 = auto()
    F2 = auto()
    F3 = auto()
    F4 = auto()
    F5 = auto()
    F6 = auto()
    S1 = 101
    S2 = 102
    S3 = 103
    S4 = 104
    S5 = 105
    S6 = 106
    S7 = 107
    S8 = 108
    S9 = 109 # 論文成長率-三年版
    

# 每一個sheet的資料
@dataclass
class Template(ABC):
    sheet_name : str
    sheet_idx : int
    _excel_func: dict[int, str] = field(default_factory=dict) 
    rows : list[list[str]] = field(default_factory=list)
    null_val_loc : set[str] = field(default_factory=set)
    _cols_without_format: set[int] = field(default_factory=set)
    dummyColsWithoutFormt : set[int] = field(default_factory=set)
    DefaultExcelFuncInteger : str = "=TEXT({0}, \"##,#0\")"
    DefaultExcelFuncPercent : str = "=TEXT({0}, \"0.00%\")"
    DefaultExcelFunc2DecimalPlace : str = "=TEXT({0}, \"0.00\")"
    DefaultExcelFuncProgress : str = "=TEXT({0}, \"+##,#0;-##,#0;##,0\")"
    _RawdataSheetSuffix :str = "rawdata"
    template_type: TemplateType | None = None # 新增的，原程式沒有

    @classmethod
    def get_rawdata_sheet_suffix(cls) -> str:
        return cls._RawdataSheetSuffix
    @classmethod
    def set_rawdata_sheet_suffix(cls, raw_data_str:str):
        cls._RawdataSheetSuffix = raw_data_str

    @property
    def sheet_tab_name(self) -> str:
        return f"sheet{self.sheet_idx}"

    @property
    def rawdata_tab_name(self) -> str:
        return f"{self.sheet_tab_name}_{self.get_rawdata_sheet_suffix()}"
    
    @property
    def cols_without_format(self) -> set[int]:
        return self._cols_without_format

    @property
    def excel_func(self) -> dict[int, str]:
        return self._excel_func

    @property
    def DefaultCellHorizontalAlignment(self) -> str:
        return "center"
    @property
    def DefaultCellVerticalAlignment(self) -> str:
        return "center"

    @abstractmethod
    def get_report(self) -> ExcelReport:
        pass

    @property
    def sort_col_idx(self) -> int:
        return 0

    # 部分報表需要排序
    def sort_rows(self, rows: list[list[str]]) -> list[list[str]]:
        sort_rows = []
        ## 判斷需要排序的表單(E1~E5)，如果沒有此需求可以移除
        if TemplateType.E1.value <= self.template_type.value <= TemplateType.E5.value:
            sort_idx = self.sort_col_idx
            if sort_idx != 0:
                sort_rows = sorted(rows, key=lambda x: float(x[sort_idx]), reverse=True)
        ## 針對領域做排序，如果沒有此需求可以移除
        elif self.template_type.value in (TemplateType.F5.value, TemplateType.F6.value):
            for s_field in ESI_FIELD_SORT:
                sort_rows.extend([x for x in rows if x[0] == s_field and len(x) > 2])

        return sort_rows if sort_rows else rows
    
    # 有排序過的報表，公式對應的idx要替換成原順序
    def get_original_row_idx(self,now_idx:int, rows: list[list[str]], check_row:list[str]) -> int :
        if TemplateType.E1.value <= self.template_type.value <= TemplateType.E5.value:
            original_r = next(o_idx  for o_idx, x in enumerate(rows) if check_row[0] == x[0])
        elif self.template_type.value in (TemplateType.F5.value, TemplateType.F6.value):
            original_r =  next(o_idx  for o_idx, x in enumerate(rows) if check_row[1] == x[1] and check_row[0] == x[0])
        else:
            original_r = now_idx
        return original_r

    def to_worksheet(self, wb : Workbook, rows: list[list[str]] ):
        self.rows = rows
        ## 建立rawdata_sheet
        self.rawdata_to_worksheet(wb, rows)

        self.rows = self.sort_rows(rows)

        report = self.get_report()
        report_rows : list[list[ECell]] = []
        for idx, rr in enumerate(self.rows):
            
            row_idx = self.get_original_row_idx(idx, rows, rr)
            rcells = self.get_cell_row(row_idx+1, rr)
            report_rows.append(rcells)

        self.rows = rows
        report.rows = report_rows
        report.export(wb, self.sheet_tab_name)   
        
    def get_cell_row(self, r:int, row : list[str]) -> list[ECell] :
        r_cells = []
        c = 1
        for col in row:
            if c in self.cols_without_format:
                r_cells.append(
                    ECell(
                        text=col, 
                        vertical_align=self.DefaultCellVerticalAlignment,
                        horizontal_align=self.DefaultCellHorizontalAlignment,
                        is_formula= False
                    )
                )
            else:
                r_cells.append(
                    ECell(
                        text=self.get_cell_value(r,c), 
                        vertical_align=self.DefaultCellVerticalAlignment,
                        horizontal_align=self.DefaultCellHorizontalAlignment,
                        is_formula= True
                    )
                )
            c+=1
        return r_cells

    def get_cell_value(self,r : int, c:int) -> str:
        cell_loc_str = self.get_cell_loc_str(r,c)
        check_str = f"{r},{c}"
        if check_str in self.null_val_loc:
            return ""
        if c in self.excel_func:
            return self.excel_func[c].format(cell_loc_str)
        else:
            return f"={cell_loc_str}"

    def get_cell_loc_str(self,r : int, c:int) -> str:
        return f"'{self.rawdata_tab_name}'!{self.get_cell_letter(r,c)}"

    def get_cell_letter(self,r : int, c:int) -> str:
        return f"{get_column_letter(c)}{r}"

    def rawdata_to_worksheet(self, wb : Workbook, rows: list[list[str]] ):
        import common.util as Util
        # 為了辨別ws是Worksheet加了: Worksheet
        ws :Worksheet = wb.create_sheet(self.rawdata_tab_name) 
        r = 1
        for row in rows:
            c = 1
            for col in row:
                cell = ws.cell(r,c)
                cell.value = Util.convert_type(col)
                if not col:
                    self.null_val_loc.add(f"{r},{c}")
                c += 1
            r+=1