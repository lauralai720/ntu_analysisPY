#!/usr/bin/env python3
"""
compare_same.py - 檢查兩個 Excel 是否完全一樣

用法:
    python compare_same.py a.xlsx b.xlsx
    python compare_same.py a.xlsx b.xlsx --tol 1e-9      (數值容許微小誤差,預設 0 = 完全相同)
    python compare_same.py a.xlsx b.xlsx --max 200       (最多印出幾筆差異,預設 50)

檢查內容(含隱藏 sheet):
  - sheet 名稱、順序、顯示/隱藏狀態
  - 每個儲存格的內容(公式比對公式文字;若檔案有快取值也一併比對)
  - 合併儲存格範圍
結果:完全一樣 → 印出「兩個檔案內容一致」,結束碼 0;有差異 → 列出差異,結束碼 1。
"""
import argparse
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def empty(v):
    return v is None or (isinstance(v, str) and v == "")


def same(a, b, tol):
    if empty(a) and empty(b):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) \
            and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= tol * max(1.0, abs(a), abs(b))
    return a == b


def main():
    ap = argparse.ArgumentParser(description="檢查兩個 Excel 是否一樣")
    ap.add_argument("file_a")
    ap.add_argument("file_b")
    ap.add_argument("--tol", type=float, default=0.0, help="數值相對容許誤差(預設 0)")
    ap.add_argument("--max", type=int, default=50, help="最多印出幾筆差異")
    a = ap.parse_args()

    fa, fb = load_workbook(a.file_a), load_workbook(a.file_b)            # 公式文字
    va, vb = load_workbook(a.file_a, data_only=True), load_workbook(a.file_b, data_only=True)  # 快取值

    diffs = []

    if fa.sheetnames != fb.sheetnames:
        only_a = [s for s in fa.sheetnames if s not in fb.sheetnames]
        only_b = [s for s in fb.sheetnames if s not in fa.sheetnames]
        if only_a:
            diffs.append(f"只有 A 有的 sheet: {only_a}")
        if only_b:
            diffs.append(f"只有 B 有的 sheet: {only_b}")
        if not only_a and not only_b:
            diffs.append(f"sheet 順序不同: A={fa.sheetnames} / B={fb.sheetnames}")

    for name in [s for s in fa.sheetnames if s in fb.sheetnames]:
        wa, wb_ = fa[name], fb[name]
        if wa.sheet_state != wb_.sheet_state:
            diffs.append(f"[{name}] 顯示狀態不同: A={wa.sheet_state} / B={wb_.sheet_state}")
        if sorted(map(str, wa.merged_cells.ranges)) != sorted(map(str, wb_.merged_cells.ranges)):
            diffs.append(f"[{name}] 合併儲存格範圍不同")
        if (wa.max_row, wa.max_column) != (wb_.max_row, wb_.max_column):
            diffs.append(f"[{name}] 範圍大小不同: A={wa.dimensions} / B={wb_.dimensions}")

        for r in range(1, max(wa.max_row, wb_.max_row) + 1):
            for c in range(1, max(wa.max_column, wb_.max_column) + 1):
                x, y = wa.cell(r, c).value, wb_.cell(r, c).value
                ref = f"{get_column_letter(c)}{r}"
                if not same(x, y, a.tol):
                    diffs.append(f"[{name}] {ref}: A={x!r} / B={y!r}")
                    continue
                cx, cy = va[name].cell(r, c).value, vb[name].cell(r, c).value
                if (cx is not None or cy is not None) and not same(cx, cy, a.tol):
                    diffs.append(f"[{name}] {ref} (快取值): A={cx!r} / B={cy!r}")

    if not diffs:
        print("兩個檔案內容一致")
        return 0

    print(f"共 {len(diffs)} 處不同(顯示前 {min(len(diffs), a.max)} 筆):")
    for d in diffs[:a.max]:
        print(" -", d)
    return 1


if __name__ == "__main__":
    sys.exit(main())